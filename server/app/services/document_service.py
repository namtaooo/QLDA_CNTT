import logging
import fitz # PyMuPDF
import docx
import openpyxl

logger = logging.getLogger(__name__)

def extract_text(filepath: str, file_type: str) -> str:
    """Extract text content from common document types."""
    content = ""
    try:
        if file_type in ("txt", "md"):
            with open(filepath, "r", encoding="utf-8") as f:
                content = f.read()
                
        elif file_type == "pdf":
            doc = fitz.open(filepath)
            for page in doc:
                content += page.get_text() + "\n"
                
        elif file_type == "docx":
            doc = docx.Document(filepath)
            content = "\n".join([para.text for para in doc.paragraphs])
            
        elif file_type == "xlsx":
            wb = openpyxl.load_workbook(filepath, data_only=True)
            for sheet in wb.worksheets:
                content += f"--- Sheet: {sheet.title} ---\n"
                for row in sheet.iter_rows(values_only=True):
                    row_data = [str(cell) for cell in row if cell is not None]
                    if row_data:
                        content += " | ".join(row_data) + "\n"
        else:
            logger.warning(f"Unsupported file type for extraction: {file_type}")
            
    except Exception as e:
        logger.error(f"Error extracting text from {filepath}: {e}")
        
    return content
